from datetime import datetime, timezone, timedelta
import logging
import os
from dags.utils.situtils.BaseDagsClient import ConfigClient
from airflow.sdk import dag, task, Asset, Metadata
import time

from dags.utils.callback import Callback
from dags.utils.models.airflow_constant import RecordEstimate, StorageEstimate, TimeEstimate
from utils.telemetry.tracing import tracer
from typing import List, Set, NamedTuple
import polars as ps
from dags.utils.postgres import PostgresEtlHook
from dags.utils.oracle import OracleEtlHook
from dags.utils.operators.ssh.ssh_operator import SSHOperatorOcbc
import requests
import json

from io import BytesIO
from pathlib import Path
import shlex
import sys
from airflow.providers.postgres.hooks.postgres import PostgresHook
from airflow.providers.ssh.hooks.ssh import SSHHook
from openpyxl import load_workbook
import pandas as pd
import xlsxwriter

_project_root = str(Path(__file__).resolve().parents[3])
if _project_root not in sys.path:
    sys.path.insert(0, _project_root)


logger = logging.getLogger(__name__)
db_control = PostgresEtlHook()
db_source = OracleEtlHook()
config = {
    "owner": "acquiring",
    "workflow_name": "dag_pten_in",
    "db_source": "oracle-db-sit-pwc",
    "db_control": "postgres-db-acquiring",
    "pten_in_file_path": "/data/airflow/nfs/artifacts/data_payment/acquiring/pten/in",
    "email_list": [],
}
FILE_PATH = os.path.dirname(os.path.realpath(__file__))
workflow_name = config["workflow_name"]
dag_name = workflow_name
callback = Callback(email_list=config["email_list"])
queue = "etl" if FILE_PATH.split("/")[-1] == "etl" else "default"
default_args = {
    "owner": config["owner"],
    "on_failure_callback": callback.task_failure_alert,
    "queue": queue,
}
tz = timezone(timedelta(hours=7))
start_date = "2026-01-01"
user_cron = None
dynamic_tags = [
    "app:payment-and-collection-platform", "job:database",
    "job:database:operation", "owner:acquiring",
    "time:between_2m_5m", "storage:below_1gb", "records:below_1k",
]


@dag(
    dag_id=workflow_name,
    schedule=user_cron,
    description="PTEN IN merchant processing and Excel generation",
    default_args=default_args,
    start_date=datetime.strptime(start_date, "%Y-%m-%d").replace(tzinfo=tz),
    catchup=False,
    dagrun_timeout=timedelta(minutes=60),
    tags=dynamic_tags,
    max_active_runs=1,
    on_success_callback=callback.dag_success_alert,
)
def dag_pten_in():
    local_file_path = config.get('pten_in_file_path')

    @task
    @tracer.start_as_current_span(f"{dag_name}_task_main_process")
    def main_process():
        today = datetime.now()
        data_changes = _read_source(
            'SELECT DISTINCT MID FROM STG_PTEN_MERCHANT_UPDATES WHERE FLAG<>:flag',
            {'flag': 'Y'},
        )
        check_data_changes = len(data_changes) != 0
        lastrun = _control_scalar(
            'select value from "PTENControlTable_TN" where "keys" = :keys',
            {'keys': 'last_run_date'},
        )
        date_last_run = datetime.fromisoformat(str(lastrun))
        if date_last_run.strftime('%Y-%m-%d') != today.strftime('%Y-%m-%d'):
            counter = 1
        else:
            counter = int(_control_scalar(
                'select value from "PTENControlTable_TN" where keys=:keys',
                {'keys': 'sequence'},
            ))
        data_new = pd.DataFrame(
            get_data_new(date_last_run, today.strftime('%Y-%m-%d %H:%M')),
            columns=[
                'NMID', 'MERCHANT_NAME', 'MPAN', 'MID', 'CITY', 'POSTAL_CODE', 'STATUS',
                'MERCHANT_CRITERIA', 'MCC', 'TOTAL_TERMINAL', 'MERCHANT_TYPE',
                'NPWP', 'KTP', 'QR_TYPE', 'MERCHANT_TYPE2', 'OPENING_DATE',
            ],
        )
        header_raw = [
            'NMID', 'MERCHANT_NAME', 'MPAN', 'MID', 'CITY', 'POSTAL_CODE', 'STATUS',
            'MERCHANT_CRITERIA', 'MCC', 'TOTAL_TERMINAL', 'MERCHANT_TYPE',
            'NPWP', 'KTP', 'QR_TYPE', 'MERCHANT_TYPE2', 'OPENING_DATE',
        ]
        try:
            logging.info('Insert data to Onboarding Logs')
            df_log_data = pd.DataFrame(data_new)
            df_log_data.columns = header_raw
            df_log_data = df_log_data[['MID', 'OPENING_DATE']]
            df_log_data['PERIOD'] = today.strftime('%Y%m%d')
            df_log_data = df_log_data[['MID', 'OPENING_DATE']]
            df_log_data['PERIOD'] = today.strftime('%Y%m%d')
            df_log_data['STATUS'] = 'PENDING'
            while True:
                logging_data = df_log_data[:5000]
                logging_data.columns = ['mid', 'opening_date', 'period', 'status']
                _insert_onboarding(
                    'QRISOnboardingLog_TH', ['mid', 'opening_date', 'period', 'status'],
                    logging_data.to_dict('records'),
                )
                logging.info('Successfully insert log data to table QRISOnboardingLog_TH')
                df_log_data = df_log_data[5000:]
                if len(df_log_data) < 1:
                    break
        except Exception as e:
            logging.error('Failed insert log data to table QRISOnboardingLog_TH: %s', e)

        filename = None
        while len(data_new) > 0:
            processed_data = data_new[:5000]
            filename, df_processed = write_data(processed_data, counter, mode='NEW')
            if filename is None:
                break
            mid = tuple(df_processed['MID'].to_list())
            _exec_control(
                'UPDATE "QRISOnboardingLog_TH" SET filename = :filename WHERE mid IN :mid',
                {'filename': filename, 'mid': mid},
            )
            data_new = data_new[5000:]
            counter += 1

        if check_data_changes:
            columns = [
                'NMID', 'MERCHANT_NAME', 'MPAN', 'MID', 'CITY', 'POSTAL_CODE', 'STATUS',
                'MERCHANT_CRITERIA', 'MCC', 'TOTAL_TERMINAL', 'MERCHANT_TYPE',
                'NPWP', 'KTP', 'QR_TYPE', 'MERCHANT_TYPE2', 'ACCEPTANCE_CHANNEL',
            ]
            data_new_update = pd.DataFrame(
                get_data_new_update(date_last_run, today.strftime('%Y-%m-%d %H:%M')),
                columns=columns,
            )
            mid_updated = tuple(data_new_update['MID'].to_list())
            _exec_control(
                'UPDATE "QRISOnboardingLog_TH" SET status=:status, updated_at=now() WHERE mid IN :mid',
                {'status': 'PENDING', 'mid': mid_updated},
            )
            while len(data_new_update) > 0:
                processed_data = data_new_update[:5000]
                filename, df_processed = write_data(processed_data, counter, mode='NEW')
                if filename is None:
                    break
                mid = tuple(df_processed['MID'].to_list())
                _exec_control(
                    'UPDATE "QRISOnboardingLog_TH" SET filename=:filename WHERE mid IN :mid',
                    {'filename': filename, 'mid': mid},
                )
                data_new_update = data_new_update[5000:]
                counter += 1

            data_update = pd.DataFrame(
                get_data_update(date_last_run, today), columns=columns,
            )
            mid_to_update = tuple(data_update['MID'].to_list())
            _exec_control(
                'UPDATE "QRISOnboardingLog_TH" SET status=:status, updated_at=now() WHERE mid IN :mid',
                {'status': 'PENDING', 'mid': mid_to_update},
            )
            while len(data_update) > 0:
                processed_data = data_update[:5000]
                filename, df_processed = write_data(processed_data, counter, mode='UPDATE')
                if filename is None:
                    break
                mid = tuple(df_processed['MID'].to_list())
                _exec_control(
                    'UPDATE "QRISOnboardingLog_TH" SET filename=:filename WHERE mid IN :mid',
                    {'filename': filename, 'mid': mid},
                )
                data_update = data_update[5000:]
                counter += 1


        _exec_control(
            'update "PTENControlTable_TN" SET value=:counter WHERE keys=:keys',
            {'counter': counter, 'keys': 'sequence'},
        )
        _exec_control(
            'update "PTENControlTable_TN" SET value=:today WHERE keys=:keys',
            {'today': today.strftime('%Y-%m-%d %H:%M'), 'keys': 'last_run_date'},
        )
        return filename

    def _read_source(sql, params=None):
        records = db_source.read_polar_with_exec(
            sql=sql, params=params, connection_id=config.get('db_source'),
        )
        df = pd.DataFrame(records.to_dicts(), columns=records.columns)
        df.columns = [column.upper() for column in df.columns]
        return df

    def _control_scalar(sql, params):
        sql = _control_sql(sql)
        result = db_control.read_polar_with_exec(
            sql=sql, params=params, connection_id=config.get('db_control'),
        )
        return result.item(0, 0)

    def _control_sql(sql):
        import re
        sql = re.sub(r'\bIN\s+:mid\b', '= ANY(:mid)', sql, flags=re.IGNORECASE)
        return re.sub(r'(?<!:):([A-Za-z_][A-Za-z0-9_]*)', r'%(\1)s', sql)

    def _exec_control(sql, params):
        params = dict(params)
        if 'mid' in params:
            params['mid'] = list(params['mid'])
            if not params['mid']:
                return
        db_control.exec_cursor(
            sql=_control_sql(sql), params=params, connection_id=config.get('db_control'),
        )

    def _insert_onboarding(table, columns, records):
        if not records:
            return
        sql = (
            f'INSERT INTO "{table}" (' + ', '.join(columns) + ') VALUES ('
            + ', '.join(f'%({column})s' for column in columns) + ')'
        )
        for record in records:
            db_control.exec_cursor(sql=sql, params=record, connection_id=config.get('db_control'))

    def get_data_new(before, now):
        logging.info('START Get Data GET_DATA_NEW PTEN IN')
        new_merchant_query = """
        SELECT
            M.EXTERNAL_ID_1 AS NMID,
            M.ACRONYM AS MERCHANT_NAME,
            MA.AP_ON_ID_1 AS MPAN,
            MA.MER_ACCEPTOR_POINT_ID AS MID,
            C.CITY_NAME AS CITY,
            M.PR_ZIP_CODE AS POSTAL_CODE,
            M.STATUS_CODE AS STATUS,
            (SELECT CARACT_VALUE FROM MER_CONTRACT_CARACTERISTICS MRC WHERE MRC.CONTRACT_ELMT_ID = MA.CONTRACT_ELMT_ID AND CARACT_CONTEXT = 'MER_CRITERIA') AS MERCHANT_CRITERIA,
            MAC.MCC AS MCC,
            (SELECT CARACT_VALUE FROM MER_CONTRACT_CARACTERISTICS MRC WHERE MRC.CONTRACT_ELMT_ID = MA.CONTRACT_ELMT_ID AND CARACT_CONTEXT = 'TOTAL_STICKER') AS TOTAL_TERMINAL,
            CASE WHEN M.LEGAL_ID_TYPE = 'P' THEN '1' ELSE '2' END AS MERCHANT_TYPE,
            MA.LEGAL_ID AS NPWP,
            MR.LEGAL_ID AS KTP,
            (SELECT CARACT_VALUE FROM MER_CONTRACT_CARACTERISTICS MRC WHERE MRC.CONTRACT_ELMT_ID = MA.CONTRACT_ELMT_ID AND CARACT_CONTEXT = 'QRIS_TYPE') AS QR_TYPE,
            (SELECT CARACT_VALUE FROM MER_CONTRACT_CARACTERISTICS MRC WHERE MRC.CONTRACT_ELMT_ID = MA.CONTRACT_ELMT_ID AND CARACT_CONTEXT = 'MERCHANT_TYPE') AS MERCHANT_TYPE2,
            MA.OPENING_DATE AS OPENING_DATE
        FROM
            MER_ACCEPTOR_POINT MA,
            MERCHANT M,
            CITY C,
            MER_ACTIVITY MAC,
            MER_REPRESENTATIVE MR,
            MER_ACCEPTANCE_CONTRACT_ELMT mace
        WHERE
            MA.MERCHANT_NUMBER = M.MERCHANT_NUMBER
            AND (M.EXTERNAL_ID_1 IS NULL AND MA.PRIVATE_DATA_1 IS NULL)
            AND (M.PR_CITY_CODE = C.CITY_CODE AND M.PR_COUNTRY_CODE = C.COUNTRY_CODE)
            AND MA.ACTIVITY_ID = MAC.ACTIVITY_ID
            AND (MA.MERCHANT_NUMBER = MR.MERCHANT_NUMBER AND MR.REPRESENTATIVE_POSITION = '01')
            AND (MA.CONTRACT_ELMT_ID = mace.CONTRACT_ELMT_ID AND mace.PRODUCT_CODE = 'IDALL00T01')
            AND MA.OPENING_DATE > to_date(:lastrun, 'YYYY-MM-DD HH24:MI')
            AND MA.OPENING_DATE < to_date(:now, 'YYYY-MM-DD HH24:MI')
        """
        logging.info('Get Data New Merchant')
        records = _read_source(new_merchant_query, {'lastrun': before.strftime('%Y-%m-%d %H:%M') if isinstance(before, datetime) else before, 'now': now.strftime('%Y-%m-%d %H:%M') if isinstance(now, datetime) else now})
        logging.info('Finish Get Data New Merchant')
        return records

    def get_data_new_update(before, now):
        logging.info('START Get Data GET_DATA_NEW_UPDATE PTEN IN')
        new_merchant_query = """
        SELECT
            M.EXTERNAL_ID_1 AS NMID,
            M.ACRONYM AS MERCHANT_NAME,
            MA.AP_ON_ID_1 AS MPAN,
            MA.MER_ACCEPTOR_POINT_ID AS MID,
            C.CITY_NAME AS CITY,
            M.PR_ZIP_CODE AS POSTAL_CODE,
            M.STATUS_CODE AS STATUS,
            (SELECT CARACT_VALUE FROM MER_CONTRACT_CARACTERISTICS MRC WHERE MRC.CONTRACT_ELMT_ID = MA.CONTRACT_ELMT_ID AND CARACT_CONTEXT = 'MER_CRITERIA') AS MERCHANT_CRITERIA,
            MAC.MCC AS MCC,
            (SELECT CARACT_VALUE FROM MER_CONTRACT_CARACTERISTICS MRC WHERE MRC.CONTRACT_ELMT_ID = MA.CONTRACT_ELMT_ID AND CARACT_CONTEXT = 'TOTAL_STICKER') AS TOTAL_TERMINAL,
            CASE WHEN M.LEGAL_ID_TYPE = 'P' THEN '1' ELSE '2' END AS MERCHANT_TYPE,
            MA.LEGAL_ID AS NPWP,
            MR.LEGAL_ID AS KTP,
            (SELECT CARACT_VALUE FROM MER_CONTRACT_CARACTERISTICS MRC WHERE MRC.CONTRACT_ELMT_ID = MA.CONTRACT_ELMT_ID AND CARACT_CONTEXT = 'QRIS_TYPE') AS QR_TYPE,
            (SELECT CARACT_VALUE FROM MER_CONTRACT_CARACTERISTICS MRC WHERE MRC.CONTRACT_ELMT_ID = MA.CONTRACT_ELMT_ID AND CARACT_CONTEXT = 'MERCHANT_TYPE') AS MERCHANT_TYPE2,
            (SELECT CARACT_VALUE FROM MER_CONTRACT_CARACTERISTICS MRC WHERE MRC.CONTRACT_ELMT_ID = MA.CONTRACT_ELMT_ID AND CARACT_CONTEXT = 'MER_ACC_CHANNEL') AS ACCEPTANCE_CHANNEL
        FROM
            MER_ACCEPTOR_POINT MA,
            MERCHANT M,
            CITY C,
            MER_ACTIVITY MAC,
            MER_REPRESENTATIVE MR,
            STG_PTEN_MERCHANT_UPDATES SPU
        WHERE
            MA.MERCHANT_NUMBER = M.MERCHANT_NUMBER
            AND (M.EXTERNAL_ID_1 IS NULL AND MA.PRIVATE_DATA_1 IS NULL)
            AND (M.PR_CITY_CODE = C.CITY_CODE AND M.PR_COUNTRY_CODE = C.COUNTRY_CODE)
            AND MA.ACTIVITY_ID = MAC.ACTIVITY_ID
            AND (MA.MERCHANT_NUMBER = MR.MERCHANT_NUMBER AND MR.REPRESENTATIVE_POSITION = '01')
            AND SPU.MID = MA.MER_ACCEPTOR_POINT_ID
            AND (SPU.FLAG <> 'Y' OR SPU.FLAG IS NULL)
        """
        logging.info('Get Data New Update Merchant')
        records = _read_source(new_merchant_query, {'lastrun': before.strftime('%Y-%m-%d %H:%M') if isinstance(before, datetime) else before, 'now': now.strftime('%Y-%m-%d %H:%M') if isinstance(now, datetime) else now})
        logging.info('Finish Get Data New Update Merchant')
        return records

    def get_data_from_nmid():
        logging.info('START FETCH MERCHANT DETAILS BY MID')
        select_mid_query = """
            SELECT DISTINCT mid
            FROM STG_PTEN_MERCHANT_UPDATES
            WHERE FLAG = 'Y'
              AND FIELD_NAME LIKE '%STATUS%'
        """
        mids = _read_source(select_mid_query)
        mid_list = mids['MID'].to_list()
        if not mid_list:
            logging.info('MID not found.')
            return []
        merchant_detail_query = """
        SELECT
            M.EXTERNAL_ID_1 AS NMID,
            M.ACRONYM AS MERCHANT_NAME,
            MA.AP_ON_ID_1 AS MPAN,
            MA.MER_ACCEPTOR_POINT_ID AS MID,
            C.CITY_NAME AS CITY,
            M.PR_ZIP_CODE AS POSTAL_CODE,
            M.STATUS_CODE AS STATUS,
            (SELECT CARACT_VALUE FROM MER_CONTRACT_CARACTERISTICS MRC WHERE MRC.CONTRACT_ELMT_ID = MA.CONTRACT_ELMT_ID AND CARACT_CONTEXT = 'MER_CRITERIA') AS MERCHANT_CRITERIA,
            MAC.MCC AS MCC,
            (SELECT CARACT_VALUE FROM MER_CONTRACT_CARACTERISTICS MRC WHERE MRC.CONTRACT_ELMT_ID = MA.CONTRACT_ELMT_ID AND CARACT_CONTEXT = 'TOTAL_STICKER') AS TOTAL_TERMINAL,
            CASE WHEN M.LEGAL_ID_TYPE = 'P' THEN '1' ELSE '2' END AS MERCHANT_TYPE,
            MA.LEGAL_ID AS NPWP,
            MR.LEGAL_ID AS KTP,
            (SELECT CARACT_VALUE FROM MER_CONTRACT_CARACTERISTICS MRC WHERE MRC.CONTRACT_ELMT_ID = MA.CONTRACT_ELMT_ID AND CARACT_CONTEXT = 'QRIS_TYPE') AS QR_TYPE,
            (SELECT CARACT_VALUE FROM MER_CONTRACT_CARACTERISTICS MRC WHERE MRC.CONTRACT_ELMT_ID = MA.CONTRACT_ELMT_ID AND CARACT_CONTEXT = 'MER_ACC_CHANNEL') AS ACCEPTANCE_CHANNEL
        FROM
            MER_ACCEPTOR_POINT MA,
            MERCHANT M,
            CITY C,
            MER_ACTIVITY MAC,
            MER_REPRESENTATIVE MR
        WHERE
            TRUNC(MA.OPENING_DATE) = TRUNC(SYSDATE)
            AND (M.EXTERNAL_ID_1 IS NOT NULL AND MA.PRIVATE_DATA_1 IS NOT NULL)
            AND MA.MERCHANT_NUMBER = M.MERCHANT_NUMBER
            AND (M.PR_CITY_CODE = C.CITY_CODE AND M.PR_COUNTRY_CODE = C.COUNTRY_CODE)
            AND MA.ACTIVITY_ID = MAC.ACTIVITY_ID
            AND (MA.MERCHANT_NUMBER = MR.MERCHANT_NUMBER AND MR.REPRESENTATIVE_POSITION = '01')
            AND M.EXTERNAL_ID_1 IN ({placeholders})
        """
        placeholders = ', '.join(f':mid_{i}' for i in range(len(mid_list)))
        merchant_detail_query = merchant_detail_query.replace('{placeholders}', placeholders)
        merchant_rows = _read_source(merchant_detail_query, {f'mid_{i}': mid for i, mid in enumerate(mid_list)})
        merchant_data = merchant_rows.to_dict('records')
        logging.info(f'Total merchant details fetched: {len(merchant_data)}')
        return merchant_data

    def get_data_update(before, now):
        logging.info('START Get Data GET_DATA_UPDATE PTEN IN')
        updated_merchant_query = """
        SELECT
            M.EXTERNAL_ID_1 AS NMID,
            M.ACRONYM AS MERCHANT_NAME,
            MA.AP_ON_ID_1 AS MPAN,
            MA.MER_ACCEPTOR_POINT_ID AS MID,
            C.CITY_NAME AS CITY,
            M.PR_ZIP_CODE AS POSTAL_CODE,
            M.STATUS_CODE AS STATUS,
            (SELECT CARACT_VALUE FROM MER_CONTRACT_CARACTERISTICS MRC WHERE MRC.CONTRACT_ELMT_ID = MA.CONTRACT_ELMT_ID AND CARACT_CONTEXT = 'MER_CRITERIA') AS MERCHANT_CRITERIA,
            MAC.MCC AS MCC,
            (SELECT CARACT_VALUE FROM MER_CONTRACT_CARACTERISTICS MRC WHERE MRC.CONTRACT_ELMT_ID = MA.CONTRACT_ELMT_ID AND CARACT_CONTEXT = 'TOTAL_STICKER') AS TOTAL_TERMINAL,
            CASE WHEN M.LEGAL_ID_TYPE = 'P' THEN '1' ELSE '2' END AS MERCHANT_TYPE,
            MA.LEGAL_ID AS NPWP,
            MR.LEGAL_ID AS KTP,
            (SELECT CARACT_VALUE FROM MER_CONTRACT_CARACTERISTICS MRC WHERE MRC.CONTRACT_ELMT_ID = MA.CONTRACT_ELMT_ID AND CARACT_CONTEXT = 'QRIS_TYPE') AS QR_TYPE,
            (SELECT CARACT_VALUE FROM MER_CONTRACT_CARACTERISTICS MRC WHERE MRC.CONTRACT_ELMT_ID = MA.CONTRACT_ELMT_ID AND CARACT_CONTEXT = 'MERCHANT_TYPE') AS MERCHANT_TYPE2,
            (SELECT CARACT_VALUE FROM MER_CONTRACT_CARACTERISTICS MRC WHERE MRC.CONTRACT_ELMT_ID = MA.CONTRACT_ELMT_ID AND CARACT_CONTEXT = 'MER_ACC_CHANNEL') AS ACCEPTANCE_CHANNEL
        FROM
            MER_ACCEPTOR_POINT MA,
            MERCHANT M,
            CITY C,
            MER_ACTIVITY MAC,
            MER_REPRESENTATIVE MR,
            STG_PTEN_MERCHANT_UPDATES SPU
        WHERE
            MA.MERCHANT_NUMBER = M.MERCHANT_NUMBER
            AND (M.EXTERNAL_ID_1 IS NOT NULL AND MA.PRIVATE_DATA_1 IS NOT NULL)
            AND (M.PR_CITY_CODE = C.CITY_CODE AND M.PR_COUNTRY_CODE = C.COUNTRY_CODE)
            AND MA.ACTIVITY_ID = MAC.ACTIVITY_ID
            AND (MA.MERCHANT_NUMBER = MR.MERCHANT_NUMBER AND MR.REPRESENTATIVE_POSITION = '01')
            AND SPU.MID = MA.MER_ACCEPTOR_POINT_ID
            AND SPU.FLAG <> 'Y'
            AND (SPU.FIELD_NAME <> 'EXTERNAL_ID_1' OR SPU.FIELD_NAME <> 'NMID')
        """
        logging.info('Get Data Update Merchant')
        records = _read_source(updated_merchant_query, {'lastrun': before.strftime('%Y-%m-%d %H:%M') if isinstance(before, datetime) else before, 'now': now.strftime('%Y-%m-%d %H:%M') if isinstance(now, datetime) else now})
        logging.info('Finish Get Data Update Merchant')
        return records

    def write_data(datas, counter, mode="NEW"):
        os.makedirs(local_file_path, exist_ok=True)
        today = datetime.now()
        strdate = today.strftime('%Y%m%d')
        col_header = [
            "Nama Merchant (max 50)", "Nama Merchant (max 25)", "MPAN", "MID",
            "Kota", "Kodepos", "Kriteria", "MCC", "Jml Terminal", "Tipe Merchant",
            "NPWP", "KTP", "Tipe QR", "Channel QR", "Tanggal Buka",
        ]
        df = pd.DataFrame(datas)
        if df.empty:
            logging.warning('write_data: empty DataFrame, skipping file write')
            return None, df
        logging.info(f'DEBUG mode={mode}: columns = {list(df.columns)}')
        df.columns = [str(column).upper() for column in df.columns]
        merchant_columns = [
            'MERCHANT_NAME', 'MERCHANT_NAME_SHORT', 'MPAN', 'MID', 'CITY',
            'POSTAL_CODE', 'MERCHANT_CRITERIA', 'MCC', 'TOTAL_TERMINAL',
            'MERCHANT_TYPE', 'NPWP', 'KTP', 'QR_TYPE', 'ACCEPTANCE_CHANNEL',
        ]
        df['MERCHANT_NAME_SHORT'] = df['MERCHANT_NAME'].fillna('').astype(str).str.slice(0, 25)
        if 'ACCEPTANCE_CHANNEL' not in df.columns:
            df['ACCEPTANCE_CHANNEL'] = df['MERCHANT_TYPE2']
        if 'OPENING_DATE' not in df.columns:
            df['OPENING_DATE'] = ''
        if mode == 'NEW':
            filename = f'QRIS_NMR_93600028_{strdate}_NEW_OCBC{counter:06}.xlsx'
            header1 = 'FORM PENDAFTARAN'
            header2 = 'NATIONAL MERCHANT REPOSITORY QRIS'
            col_header.extend(['Interface', 'Model Usaha'])
            df['INTERFACE'] = '1'
            df['MODEL_USAHA'] = 'model_usaha'
            df = df[merchant_columns + ['OPENING_DATE', 'INTERFACE', 'MODEL_USAHA']]
        elif mode == 'UPDATE':
            filename = f'UPDATE_QRIS_NMR_93600028_{strdate}_NEW_OCBC{counter:06}.xlsx'
            header1 = 'FORM UPDATE'
            header2 = 'NATIONAL MERCHANT REPOSITORY QRIS'
            col_header = col_header[:-1] + ['Interface', 'Model Usaha', 'NMID', 'Keterangan UPDATE']
            df['INTERFACE'] = '1'
            df['MODEL_USAHA'] = 'model_usaha'
            df['UPDATE_REASON'] = 'Data Merchant Update'
            df = df[merchant_columns + ['INTERFACE', 'MODEL_USAHA', 'NMID', 'UPDATE_REASON']]
        elif mode == 'DELETE':
            filename = f'DELETE_QRIS_NMR_93600028_{strdate}_NEW_OCBC{counter:06}.xlsx'
            header1 = 'FORM PENGHAPUSAN'
            header2 = 'NATIONAL MERCHANT REPOSITORY QRIS'
            df = df[merchant_columns + ['OPENING_DATE']]
        else:
            raise ValueError(f'Unsupported PTEN mode: {mode}')
        df = df.astype(object).where(pd.notna(df), '')
        df.columns = col_header
        df.insert(0, 'No', range(1, len(df) + 1))
        full_path = os.path.join(local_file_path, filename)
        wb = xlsxwriter.Workbook(full_path)
        ws = wb.add_worksheet('Sheet1')
        ws.write(0, 0, header1)
        ws.write(1, 0, header2)
        for col_num, col_name in enumerate(df.columns):
            ws.write(2, col_num, col_name)
        for row_idx, row_data in enumerate(df.itertuples(index=False)):
            for col_num, value in enumerate(row_data):
                ws.write(3 + row_idx, col_num, value)
        ws.write(len(df) + 3, 0, 'end')
        wb.close()
        return full_path, df

    @task
    @tracer.start_as_current_span(f"{dag_name}_task_update_status_log")
    def update_status_log():
        db_source.exec_cursor(
            'UPDATE STG_PTEN_MERCHANT_UPDATES SET FLAG=:flag WHERE FLAG<>:flag OR FLAG IS NULL',
            params={'flag': 'Y'}, connection_id=config.get('db_source'),
        )
        _exec_control(
            'update "QRISOnboardingLog_TH" SET status=:new_status, '
            'last_attempt_at=current_timestamp, attempts=attempts+1, updated_at=now() '
            'WHERE status=:status',
            {'new_status': 'SENT', 'status': 'PENDING'},
        )
        logging.info('Successfully update status onboarding to SENT')

    process = main_process()
    status = update_status_log()
    process >> status


dag_pten_in()
